import io as io_module
import re
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from python_calamine import CalamineWorkbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from sql import build_in_clause

from services.processor import CUTOVER_DAY


# ─── Нормативы (минуты) ───
# done_day <  CUTOVER_DAY → carte (старый справочник): absolute / planned / detail.
# done_day >= CUTOVER_DAY → work_catalog (новый справочник): norm / planned / detail (short=task_report).

MKD_OBJECT_TYPES = ("Квартира", "Коммунальная квартира", "Дом блокированной застройки")
MKD_IN = ", ".join(f"'{t}'" for t in MKD_OBJECT_TYPES)
BP_MKD = 40   # «Выявление безучетного потребления БП МКД»
BP_IZHS = 60  # «Выявление безучетного потребления БП ИЖС»

ALCOR_TITLE = "Выполнение задания в Алькоре"


_IN_CHUNK = 32500


# ─── 1С-файл (вкладка «1С» отчёта по нормативам) ───
# TDSheet («Полная выгрузка из 1С УЛР»): «Дата работ» (13/N) — период строки,
# «Поручение контрагента» (10/K), «Непосредственный исполнитель1/2» (22/23 W/X),
# «Заказчик.Входит в группу» (7/H, ЛЭ ≡ РЛЭ / ПСК), «Проект строка» (5/F).
COL_1C_DATE = 13
COL_1C_WORK = 10
COL_1C_EXECUTOR1 = 22
COL_1C_EXECUTOR2 = 23
COL_1C_PROJECT = 5
COL_1C_CUSTOMER = 7

# «Дата работ» > 30.09.2026 → норматив из work_catalog, иначе из carte.
CUTOVER_1C = (2026, 9, 30)

# Маппинг «Поручение контрагента» (для «Снятие показаний» ещё и «Проект строка») →
# work_catalog.short по заказчику. Файловая «ЛЭ» = «РЛЭ» (алиасы).
# Пустая ячейка (None) → нет маппинга: выводим строку с названием 1С и нормативом 0.
WORK_MAP_1C_WC = {
    ("Снятие показаний", "Текущие"): {"РЛЭ": "Штучное КСП", "ПСК": "Штучное КСП ЮЛ"},
    ("Снятие показаний", "Плановые"): {"РЛЭ": "Массовое КСП", "ПСК": "Плановое КСП ЮЛ"},
    ("Снятие показаний", 'Зона ООО "РКС-энерго"'): {"РЛЭ": "Штучное КСП", "ПСК": None},
    ("Инструментальная проверка", None): {"РЛЭ": "Инструментальная проверка 3ф", "ПСК": "ИП ЮЛ"},
    ("Периодический контроль ЮЛ", None): {"РЛЭ": "Проверка РО РЛЭ", "ПСК": "Проверка РО ЮЛ"},
    ("Осмотр", None): {"РЛЭ": None, "ПСК": "Осмотры ЮЛ"},
    ("Инструментальная проверка 3ф", None): {"РЛЭ": "Инструментальная проверка 3ф", "ПСК": "ИП ЮЛ"},
    ("Допуск ПУ БП", None): {"РЛЭ": "Допуск ПУ 1ф", "ПСК": "Допуск ПУ ФЛ"},
    ("Допуск ПУ", None): {"РЛЭ": "Допуск ПУ 3ф", "ПСК": "Допуск ПУ ЮЛ"},
    ("Замеры", None): {"РЛЭ": "Чтение архива РЛЭ", "ПСК": "Замеры"},
    ("Инструментальная проверка 1ф", None): {"РЛЭ": "Инструментальная проверка 1ф", "ПСК": "ИП ФЛ"},
    ("Акт о недопуске", None): {"РЛЭ": "Недопуск без уведомления", "ПСК": "Недопуск штучное КСП ЮЛ"},
    ("Архив", None): {"РЛЭ": "Чтение архива РЛЭ", "ПСК": "Чтение архива ПСК"},
    ("Допуск ПУ 1ф", None): {"РЛЭ": "Допуск ПУ 1ф", "ПСК": None},
    ("Допуск ПУ ИЖС", None): {"РЛЭ": "Допуск ПУ 3ф", "ПСК": None},
    ("Допуск ПУ 3ф", None): {"РЛЭ": "Допуск ПУ 3ф", "ПСК": "Допуск ПУ ЮЛ"},
    ("Бытовая заявка", None): {"РЛЭ": "Штучное КСП", "ПСК": "Штучное КСП ФЛ"},
    ("Доставка уведомлений", None): {"РЛЭ": None, "ПСК": "Уведомление без подписи"},
    ("Уведомление под подпись", None): {"РЛЭ": None, "ПСК": "Уведомление под подпись"},
    ("Уведомление без подписи", None): {"РЛЭ": None, "ПСК": "Уведомление без подписи"},
    ("Неучтенное потребление ЮЛ", None): {"РЛЭ": "Выявление безучетного потребления ЮЛ", "ПСК": "Выявление безучетного потребления ЮЛ"},
    ("Жалоба на некач. эл.", None): {"РЛЭ": None, "ПСК": "Жалобы на качество"},
}

SKIP_1C_WORKS = {"Возврат на доработку"}


def _norm_name(name):
    """Нормализация ФИО «ё/е» для сравнения (как в rle._norm)."""
    return (name or "").strip().replace("ё", "е").replace("Ё", "Е")


def _cell(row, idx) -> str:
    """Значение ячейки как строка (пустая, если нет/пусто)."""
    if idx < len(row) and row[idx] is not None and row[idx] != "":
        return str(row[idx]).strip()
    return ""


def _parse_date(value):
    """«Дата работ» → (год, месяц, день) или None. Понимает dd.mm.yyyy и datetime."""
    if value is None or value == "":
        return None
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        return (value.year, value.month, value.day)
    s = str(value).strip()
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$", s)
    if m:
        return (int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return (int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def _as_int(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _read_1c_sheet(content: bytes) -> list[list]:
    """Читает TDSheet из 1С-файла (calamine, fallback openpyxl)."""
    try:
        wb = CalamineWorkbook.from_filelike(io_module.BytesIO(content))
        return [list(row) for row in wb.get_sheet_by_name("TDSheet").to_python()]
    except Exception:
        wb = load_workbook(io_module.BytesIO(content), read_only=True, data_only=True)
        return [list(row) for row in wb["TDSheet"].iter_rows(values_only=True)]


async def build_1c_rows(db_session: AsyncSession, content: bytes, period: str = "") -> list[list]:
    """Строки вкладки «1С»: [ФИО, Должность, Отделение, Работа, Количество, Норматив].

    Берём строки, у которых месяц «Дата работ» совпадает с выбранным периодом.
    Норматив: «Дата работ» <= 30.09.2026 → carte (detail → title/absolute),
    иначе → work_catalog (WORK_MAP_1C_WC по заказчику/проекту, norm из work_catalog).
    «Возврат на доработку» пропускаем. Исполнитель — «Непосредственный исполнитель1/2».
    ФИО/должность/отделение — из users по executor_name (как во вкладках «Алькор»/«Проект»).
    """
    sheet = _read_1c_sheet(content)

    carte_map: dict[str, tuple[str, int]] = {}
    carte_rows = (await db_session.execute(text(
        "SELECT title, absolute, detail FROM carte WHERE detail IS NOT NULL AND detail != ''"
    ))).fetchall()
    for title, absolute, detail in carte_rows:
        for part in str(detail).split("ИЛИ"):
            part = part.strip()
            if part:
                carte_map[part] = (title or "", _as_int(absolute) or 0)

    wc_norm = {}
    wc_rows = (await db_session.execute(text(
        "SELECT short, norm FROM work_catalog WHERE norm IS NOT NULL"
    ))).fetchall()
    for short, norm in wc_rows:
        wc_norm[short] = _as_int(norm)

    users = (await db_session.execute(text(
        "SELECT full_name, executor_name, position, dept FROM users"
    ))).fetchall()
    user_by_exec = {}
    for full_name, executor_name, position, dept in users:
        if executor_name:
            user_by_exec[_norm_name(executor_name)] = (full_name or "", position or "", dept or "")

    counts: dict[tuple[str, str], int] = defaultdict(int)
    norm_by_work: dict[str, int] = {}
    for row in sheet[1:]:
        if not row:
            continue
        parsed = _parse_date(row[COL_1C_DATE] if len(row) > COL_1C_DATE else None)
        if parsed is None:
            continue
        y, m, d = parsed
        if f"{y} {m:02d}" != period:
            continue

        work = _cell(row, COL_1C_WORK)
        if not work or work in SKIP_1C_WORKS:
            continue

        if (y, m, d) > CUTOVER_1C:
            customer = _cell(row, COL_1C_CUSTOMER)
            customer = "РЛЭ" if customer == "ЛЭ" else customer  # «ЛЭ» ≡ «РЛЭ»
            project = _cell(row, COL_1C_PROJECT) if work == "Снятие показаний" else None
            entry = WORK_MAP_1C_WC.get((work, project))
            short = entry.get(customer) if entry else None
            if short and short in wc_norm:
                title, norm = short, wc_norm[short]
            else:
                title, norm = work, 0  # нет маппинга → норматив 0
        else:
            title, norm = carte_map.get(work, (work, 0))

        norm_by_work[title] = norm
        for idx in (COL_1C_EXECUTOR1, COL_1C_EXECUTOR2):
            executor = _cell(row, idx)
            if executor:
                counts[(executor, title)] += 1

    if not counts:
        return []

    rows = []
    for (executor, title), count in counts.items():
        full_name, position, dept = user_by_exec.get(_norm_name(executor), (executor, "", ""))
        rows.append([full_name, position, dept, title, count, count * norm_by_work[title]])
    rows.sort(key=lambda r: (_norm_name(r[0]), r[3]))
    return rows


async def apply_norms(db_session: AsyncSession, task_numbers: list[str] | None = None) -> None:
    """Проставляет main_afl.norm (база) и main_afl.extra (доп.).

    done_day <  CUTOVER_DAY — по старому справочнику carte;
    done_day >= CUTOVER_DAY — по новому справочнику work_catalog.
    norm — основной вид работ; extra — остальное («Код …», «Причина…», «+5 Алькор»).
    «Дубли» и «Ручная проверка» → norm=0, extra=0.
    """
    if task_numbers:
        # Бьём на чанки, чтобы не упереться в лимит SQLite по числу bind-параметров.
        for i in range(0, len(task_numbers), _IN_CHUNK):
            await _apply_norms_scoped(db_session, task_numbers[i:i + _IN_CHUNK])
    else:
        await _apply_norms_scoped(db_session, None)


async def _apply_norms_scoped(db_session: AsyncSession, task_numbers: list[str] | None) -> None:
    scope = ""
    scope_params: dict = {}
    if task_numbers:
        names, bind = build_in_clause("n", task_numbers)
        scope = f" AND task_number IN ({names})"
        scope_params.update(bind)

    async def run(set_clause: str, condition: str) -> None:
        await db_session.execute(
            text(f"UPDATE main_afl SET {set_clause} WHERE {condition}{scope}"),
            dict(scope_params),
        )

    old = f"done_day < '{CUTOVER_DAY}'"     # старый справочник carte
    new = f"done_day >= '{CUTOVER_DAY}'"    # новый справочник work_catalog

    # Сброс
    await run("norm = NULL, extra = NULL", "1=1")

    # 0. Дубли / Ручная проверка → 0
    await run("norm = 0, extra = 0", "task_detail IN ('Дубли', 'Ручная проверка')")

    # 1. Базовый тариф (norm)
    await run(
        "norm = (SELECT absolute FROM carte WHERE carte.kind = 'base' AND carte.title = main_afl.task_report LIMIT 1), extra = 0",
        f"task_report IS NOT NULL AND COALESCE(task_detail, '') NOT IN ('Дубли', 'Ручная проверка') AND {old}",
    )
    await run(
        "norm = (SELECT norm FROM work_catalog WHERE work_catalog.kind = 'base' AND work_catalog.short = main_afl.task_report LIMIT 1), extra = 0",
        f"task_report IS NOT NULL AND COALESCE(task_detail, '') NOT IN ('Дубли', 'Ручная проверка') AND {new}",
    )

    # 1a. «Перепрограммирование ПУ»: вид работ остаётся «Бытовые заявки», но норматив 50 (старый справочник).
    await run(
        "norm = 50",
        f"work_type_in_task = 'Перепрограммирование ПУ' AND task_report = 'Бытовые заявки' AND COALESCE(task_detail, '') NOT IN ('Дубли', 'Ручная проверка') AND {old}",
    )

    # 2. «Акт неучтённого потребления» → замещающий норматив «Выявление безучетного потребления»
    # (не вид работ, влияет только на нормативы): ИЖС = 60, МКД = 40.
    # Оверрайд только если оригинальный норматив не выше безучётного.
    await run(
        f"norm = CASE WHEN service_object_type IN ({MKD_IN}) THEN {BP_MKD} ELSE {BP_IZHS} END, extra = 0",
        f"task_detail = 'Акт неучтённого потребления' "
        f"AND COALESCE(norm, 0) <= CASE WHEN service_object_type IN ({MKD_IN}) THEN {BP_MKD} ELSE {BP_IZHS} END",
    )

    # 3. Замещающие тарифы: norm=0, extra=тариф
    await run(
        "norm = 0, extra = (SELECT absolute FROM carte WHERE carte.kind = 'replacement' AND carte.detail = main_afl.task_detail LIMIT 1)",
        f"task_detail IN (SELECT detail FROM carte WHERE carte.kind = 'replacement' AND carte.detail IS NOT NULL) AND COALESCE(task_type, '') != 'Плановый' AND {old}",
    )
    await run(
        "norm = 0, extra = (SELECT planned FROM carte WHERE carte.kind = 'replacement' AND carte.detail = main_afl.task_detail LIMIT 1)",
        f"task_detail IN (SELECT detail FROM carte WHERE carte.kind = 'replacement' AND carte.detail IS NOT NULL) AND task_type = 'Плановый' AND {old}",
    )
    await run(
        "norm = 0, extra = (SELECT norm FROM work_catalog WHERE work_catalog.kind = 'replacement' AND work_catalog.detail = main_afl.task_detail LIMIT 1)",
        f"task_detail IN (SELECT detail FROM work_catalog WHERE work_catalog.kind = 'replacement' AND work_catalog.detail IS NOT NULL) AND COALESCE(task_type, '') != 'Плановый' AND {new}",
    )
    await run(
        "norm = 0, extra = (SELECT planned FROM work_catalog WHERE work_catalog.kind = 'replacement' AND work_catalog.detail = main_afl.task_detail LIMIT 1)",
        f"task_detail IN (SELECT detail FROM work_catalog WHERE work_catalog.kind = 'replacement' AND work_catalog.detail IS NOT NULL) AND task_type = 'Плановый' AND {new}",
    )

    # 4. Дополнительные тарифы: extra += тариф
    await run(
        "extra = extra + (SELECT absolute FROM carte WHERE carte.kind = 'additional' AND carte.detail = main_afl.task_detail LIMIT 1)",
        f"task_detail IN (SELECT detail FROM carte WHERE carte.kind = 'additional' AND carte.detail IS NOT NULL) AND {old}",
    )
    await run(
        "extra = extra + (SELECT norm FROM work_catalog WHERE work_catalog.kind = 'additional' AND work_catalog.detail = main_afl.task_detail LIMIT 1)",
        f"task_detail IN (SELECT detail FROM work_catalog WHERE work_catalog.kind = 'additional' AND work_catalog.detail IS NOT NULL) AND {new}",
    )

    # 5. «Выполнение задания в Алькоре» +5 → extra
    # (дублям по task_report не начисляем, даже если их task_detail был позже перезаписан)
    await run(
        f"extra = extra + (SELECT absolute FROM carte WHERE carte.title = '{ALCOR_TITLE}' LIMIT 1)",
        f"COALESCE(task_report, '') <> 'Дубли' AND COALESCE(task_detail, '') NOT IN ('Дубли', 'Ручная проверка') AND (norm IS NOT NULL OR extra IS NOT NULL) AND {old}",
    )
    await run(
        f"extra = extra + (SELECT norm FROM work_catalog WHERE work_catalog.short = '{ALCOR_TITLE}' LIMIT 1)",
        f"COALESCE(task_report, '') <> 'Дубли' AND COALESCE(task_detail, '') NOT IN ('Дубли', 'Ручная проверка') AND (norm IS NOT NULL OR extra IS NOT NULL) AND {new}",
    )

    # 6. Без вида работ и без замещающего/дополнительного тарифа → norm=0, extra=0
    # (иначе остаётся NULL — как у «невозможно без причины», task_detail='0').
    await run("norm = 0, extra = 0", "norm IS NULL AND extra IS NULL")


async def apply_manual_norm(db_session: AsyncSession, task_numbers: list[str]) -> None:
    """Норматив после ручной смены вида работ: norm = базовый тариф, extra = «+5 Алькор».

    При пустом `task_report` (смена на «без вида работ») норматив не начисляем:
    norm=0, extra=0 — иначе строка получала бы +5 «Выполнение задания в Алькоре»
    без основного вида работ.
    """
    names, params = build_in_clause("n", task_numbers)
    await db_session.execute(text(f"""
        UPDATE main_afl SET
            norm = CASE
                WHEN task_report IS NULL THEN 0
                WHEN done_day < '{CUTOVER_DAY}' THEN (SELECT absolute FROM carte WHERE carte.kind = 'base' AND carte.title = main_afl.task_report LIMIT 1)
                WHEN done_day >= '{CUTOVER_DAY}' THEN (SELECT norm FROM work_catalog WHERE work_catalog.short = main_afl.task_report LIMIT 1)
                ELSE 0
            END,
            extra = CASE
                WHEN task_report IS NULL THEN 0
                WHEN done_day < '{CUTOVER_DAY}' THEN (SELECT absolute FROM carte WHERE carte.title = '{ALCOR_TITLE}' LIMIT 1)
                WHEN done_day >= '{CUTOVER_DAY}' THEN (SELECT norm FROM work_catalog WHERE work_catalog.short = '{ALCOR_TITLE}' LIMIT 1)
                ELSE 0
            END
        WHERE task_number IN ({names})
    """), params)


async def aggregate_utalo(db_session: AsyncSession) -> int:
    """Агрегирует все main_afl (norm != 0 или extra != 0) в utalo.

    База: task_report = основной вид работ, norm_sum = norm.
    Доп.: task_report = источник extra («Код …» или «Выполнение задания в Алькоре»), norm_sum = extra.
    """
    await db_session.execute(text("DELETE FROM utalo"))
    await db_session.execute(text(f"""
        INSERT INTO utalo (period, staff_id, full_name, position, dept, task_report, count, norm_sum)
        SELECT SUBSTR(m.done_day, 1, 4) || ' ' || SUBSTR(m.done_day, 6, 2),
               u.staff_id, u.full_name, u.position, u.dept, m.task_report,
               COUNT(*), SUM(m.norm)
        FROM main_afl m
        JOIN users u ON u.executor_name = m.executor
        WHERE m.norm != 0 AND m.done_day IS NOT NULL
        GROUP BY 1, 2, 3, 4, 5, 6

        UNION ALL

        SELECT SUBSTR(m.done_day, 1, 4) || ' ' || SUBSTR(m.done_day, 6, 2),
               u.staff_id, u.full_name, u.position, u.dept,
               COALESCE(
                   (SELECT c.title FROM carte c WHERE c.detail = m.task_detail AND c.detail IS NOT NULL LIMIT 1),
                   'Выполнение задания в Алькоре'
               ),
               COUNT(*), SUM(m.extra)
        FROM main_afl m
        JOIN users u ON u.executor_name = m.executor
        WHERE m.extra != 0 AND m.done_day IS NOT NULL
        GROUP BY 1, 2, 3, 4, 5, 6
    """))
    await db_session.commit()
    return (await db_session.execute(text("SELECT COUNT(*) FROM utalo"))).scalar()


def generate_premium_xlsx_bytes(line_rows, s1c_rows, alcor_rows, project_rows) -> bytes:
    """Xlsx отчёта по нормативам: вкладки «Линия», «1С», «Алькор» (база) и «Проект» (доп.)."""
    wb = Workbook()

    ws_line = wb.active
    ws_line.title = "Линия"
    ws_line.append(["Табель", "ФИО", "Должность", "Отделение", "Время"])
    for row in line_rows:
        ws_line.append(list(row))

    headers = ["ФИО", "Должность", "Отделение", "Работа", "Количество", "Норматив"]

    ws_1c = wb.create_sheet("1С")
    ws_1c.append(headers)
    for row in s1c_rows:
        ws_1c.append(list(row))

    ws_alcor = wb.create_sheet("Алькор")
    ws_project = wb.create_sheet("Проект")

    for ws, rows in ((ws_alcor, alcor_rows), (ws_project, project_rows)):
        ws.append(headers)
        for row in rows:
            ws.append(list(row))

    output = io_module.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
