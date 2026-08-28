import io

from openpyxl import load_workbook
from python_calamine import CalamineWorkbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from data.models import Base
from data.config import engine

RAW_AFL_XLSX_COLUMNS = [
    'task_number', 'task_source', 'task_type', 'work_type_in_task',
    'created_at', 'address', 'municipal_district', 'house_type',
    'personal_account', 'contract_status', 'contract_created_at',
    'debt_amount', 'debt_calculation_date',
    'disconnection_reconnection_debt_amount', 'debt_calculation_date_1',
    'notification_debt_amount', 'debt_calculation_date_2', 'due_date',
    'service_object_type', 'subscriber_name', 'subscriber_type',
    'subscriber_phone', 'metering_point', 'meter_type', 'meter_model',
    'meter_serial_number', 'manufacture_year', 'last_verification_date',
    'meter_tariff_rate', 'integer_capacity', 'fractional_capacity',
    'calibration_interval', 'rated_current', 'rated_voltage',
    'meter_installation_place', 'meter_status', 'meter_ownership',
    'active_disconnection', 'last_readings_t1_estimated',
    'last_readings_t2_estimated', 'last_readings_t3_estimated',
    'last_estimated_readings_date', 'last_readings_t1_control',
    'last_readings_t2_control', 'last_readings_t3_control',
    'last_control_readings_date', 'seals', 'has_current_transformer',
    'has_voltage_transformer', 'meter_inspection_results',
    'meter_type_1', 'meter_model_1', 'meter_serial_number_1',
    'manufacture_year_1', 'last_verification_date_1', 'meter_tariff_rate_1',
    'integer_capacity_1', 'fractional_capacity_1', 'calibration_interval_1',
    'rated_current_1', 'rated_voltage_1', 'meter_installation_place_1',
    'meter_ownership_1', 'seals_1', 't1', 't2', 't3', 'violations',
    'meter_status_date', 'meter_malfunction', 'unauthorized_interference',
    'unauthorized_connection', 'additional_violations',
    'unsuccessful_inspection_reason', 'work_type', 'work_result',
    'meter_type_2', 'meter_model_2', 'meter_serial_number_2',
    'manufacture_year_2', 'last_verification_date_2', 'meter_tariff_rate_2',
    'integer_capacity_2', 'fractional_capacity_2', 'calibration_interval_2',
    'rated_current_2', 'rated_voltage_2', 'meter_installation_place_2',
    'meter_ownership_2', 'seals_2', 't1_1', 't2_1', 't3_1',
    'final_meter_status', 'acts', 'comment', 'work_start_date',
    'work_end_date', 'sent_to_billing', 'billing_sent_at',
    'task_organization', 'third_party_organization', 'assignee',
    'executor', 'executor_organization', 'visit_reason', 'verified',
    'status', 'work_result_1', 'status_changed_at', 'task_link',
]


def _read_excel_rows(content: bytes) -> list[list]:
    """Читает xlsx → список строк (список списков значений). calamine (Rust) → fallback openpyxl."""
    try:
        wb = CalamineWorkbook.from_filelike(io.BytesIO(content))
        sheet = wb.get_sheet_by_index(0)
        return [list(row) for row in sheet.to_python()]
    except Exception:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]


def _cell_to_str(value):
    """Повторяет pandas dtype=str: пустая ячейка/'' → None, целочисленный float → '123', иначе str."""
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


async def load_xlsx_to_raw(db_session: AsyncSession, content: bytes) -> tuple[bool, str, int]:
    try:
        rows = _read_excel_rows(content)
        if not rows:
            return False, "Файл пуст", 0

        header = rows[0]
        data = rows[1:]

        # Если колонок больше, чем ожидаем, отбрасываем первую (служебный индекс)
        if len(header) > len(RAW_AFL_XLSX_COLUMNS):
            header = header[1:]
            data = [row[1:] for row in data]

        if len(header) != len(RAW_AFL_XLSX_COLUMNS):
            return False, f"Неверное количество колонок: {len(header)} вместо {len(RAW_AFL_XLSX_COLUMNS)}", 0

        total_rows = len(data)
        if total_rows == 0:
            return False, "Файл пуст", 0

        records = []
        for row in data:
            rec = {}
            for i, name in enumerate(RAW_AFL_XLSX_COLUMNS):
                rec[name] = _cell_to_str(row[i]) if i < len(row) else None
            records.append(rec)

        # Удаляем старую таблицу и создаём новую
        await db_session.execute(text("DROP TABLE IF EXISTS raw_afl"))
        await db_session.commit()

        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)

        columns = ", ".join(f'"{c}"' for c in RAW_AFL_XLSX_COLUMNS)
        placeholders = ", ".join(f":{c}" for c in RAW_AFL_XLSX_COLUMNS)
        insert_sql = f"INSERT INTO raw_afl ({columns}) VALUES ({placeholders})"
        async with engine.begin() as conn:
            await conn.execute(text(insert_sql), records)

        return True, "", total_rows

    except Exception as e:
        return False, f"Ошибка загрузки: {str(e)}", 0
